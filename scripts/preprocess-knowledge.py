"""Read-only source conversion. No model, network, or source-code execution.

This maintenance tool uses the bundled document runtime. The desktop app has its
own native importers; users do not need Python or this script.
"""
from pathlib import Path
import argparse, ast, hashlib, json, re, csv, io, zipfile
from datetime import datetime, timezone

TEXT = {'.txt','.md','.csv','.json','.yaml','.yml','.xml','.html','.htm','.rtf',
        '.py','.cpp','.c','.h','.hpp','.rs','.js','.jsx','.ts','.tsx','.java','.sh',
        '.ps1','.launch','.rviz','.world','.typ','.toml','.ini','.cfg','.pos'}
SKIP = {'.git','node_modules','__pycache__','.venv','venv','target','dist','tmp'}

def decode(data):
    for encoding in ('utf-8-sig','utf-16','gb18030'):
        try: return data.decode(encoding)
        except UnicodeError: pass
    raise ValueError('无法确定文字编码，未用乱码替代')

def convert(path):
    ext=path.suffix.lower(); warnings=[]
    if ext in TEXT:
        text=decode(path.read_bytes())
        if ext == '.csv':
            rows=csv.reader(io.StringIO(text)); samples=[]; count=0
            for row in rows:
                count+=1
                if count<=21: samples.append('\t'.join(row))
            warnings.append(f'数据表共 {count} 行，仅提取字段与前 20 条样例；样例不能代表统计结论')
            return f'[数据表结构与样例 · 总行数 {count}]\n'+'\n'.join(samples),warnings
        if ext == '.py':
            try:
                tree=ast.parse(text)
                landmarks={node.lineno: f'函数/类 {node.name}' for node in ast.walk(tree)
                           if isinstance(node,(ast.FunctionDef,ast.AsyncFunctionDef,ast.ClassDef))}
            except SyntaxError: landmarks={};warnings.append('Python 源码语法未解析，仍完整保存原文')
        else: landmarks={}
        code=ext not in {'.txt','.md','.csv','.html','.htm','.rtf'}
        if code:
            lines=[]
            for number,line in enumerate(text.splitlines(),1):
                if number in landmarks: lines.append(f'\n[实现单元：{landmarks[number]}，原始行 {number}]')
                lines.append(f'{number:05d} | {line}')
            text='\n'.join(lines)
        return text,warnings
    if ext=='.docx':
        from docx import Document
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        try: document=Document(path)
        except Exception:
            from xml.etree import ElementTree as ET
            with zipfile.ZipFile(path) as archive: tree=ET.fromstring(archive.read('word/document.xml'))
            parts=[]
            for i,p in enumerate(tree.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'),1):
                text=''.join(t.text or '' for t in p.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
                if text.strip(): parts.append(f'[段落 {i}]\n{text}')
            warnings.append('内部链接异常，改为直接读取文档 XML 文字层；表格布局与附件未解析')
            return '\n\n'.join(parts),warnings
        parts=[]; section=0
        for child in document.element.body:
            if child.tag.endswith('}p'):
                paragraph=Paragraph(child,document);section+=1
                if paragraph.text.strip(): parts.append(f'[段落 {section} · {paragraph.style.name}]\n{paragraph.text}')
            elif child.tag.endswith('}tbl'):
                table=Table(child,document)
                parts.append('[表格]\n'+'\n'.join('\t'.join(cell.text for cell in row.cells) for row in table.rows))
        return '\n\n'.join(parts),warnings
    if ext=='.pptx':
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        def shapes_text(shapes):
            parts=[]
            for shape in shapes:
                if shape.shape_type==MSO_SHAPE_TYPE.GROUP: parts.extend(shapes_text(shape.shapes))
                elif shape.has_table: parts.append('[表格]\n'+'\n'.join('\t'.join(c.text for c in row.cells) for row in shape.table.rows))
                elif shape.has_text_frame and shape.text.strip(): parts.append(shape.text)
            return parts
        parts=[]
        for i,slide in enumerate(Presentation(path).slides,1):
            text='\n'.join(shapes_text(slide.shapes))
            if slide.has_notes_slide:
                notes=slide.notes_slide.notes_text_frame.text
                if notes.strip(): text+='\n[演讲备注]\n'+notes
            parts.append(f'[幻灯片 {i}]\n{text}')
        warnings.append('图像、图表中的文字和空间关系未作视觉识别')
        return '\n\n'.join(parts),warnings
    if ext=='.pdf':
        import fitz
        parts=[]
        with fitz.open(path) as pdf:
            for i,page in enumerate(pdf,1):
                text=page.get_text('text',sort=True)
                if not text.strip(): warnings.append(f'第 {i} 页没有文字层，需要 OCR/视觉识别')
                parts.append(f'[PDF 第 {i} 页]\n{text}')
        return '\n\n'.join(parts),warnings
    if ext=='.xlsx':
        from openpyxl import load_workbook
        values=load_workbook(path,read_only=True,data_only=True)
        formulas=load_workbook(path,read_only=True,data_only=False)
        parts=[]
        for sheet in values:
            rows=[]
            for row,formula_row in zip(sheet.iter_rows(),formulas[sheet.title].iter_rows()):
                cells=[]
                for cell,formula_cell in zip(row,formula_row):
                    value=cell.value
                    if value is None:
                        formula=formula_cell.value
                        if isinstance(formula,str) and formula.startswith('='): value=f'[公式，未计算] {formula}'
                    if value is not None: cells.append(f'{cell.coordinate}: {value}')
                if cells: rows.append('\t'.join(cells))
            parts.append(f'[工作表 {sheet.title}]\n'+'\n'.join(rows))
        values.close();formulas.close()
        return '\n\n'.join(parts),warnings
    raise ValueError('不在文字预处理范围：保留原文件；图片/视频需另行识别，权重/压缩包不作为事实文字')

def main():
    parser=argparse.ArgumentParser();parser.add_argument('source');parser.add_argument('output');args=parser.parse_args()
    source=Path(args.source).resolve();output=Path(args.output).resolve()
    if not source.is_dir() or output==source or source in output.parents:
        raise ValueError('输出必须是原资料目录之外的独立目录')
    output.mkdir(parents=True,exist_ok=True);records=[];existing={}
    if (output/'manifest.json').exists():
        existing={r['source']:r for r in json.loads((output/'manifest.json').read_text(encoding='utf-8'))['records']}
    for path in sorted(source.rglob('*')):
        if not path.is_file(): continue
        relative=path.relative_to(source)
        if any(part in SKIP for part in relative.parts) or path.name.startswith(('~$','.env')): continue
        hasher=hashlib.sha256()
        with path.open('rb') as stream:
            for data in iter(lambda:stream.read(1024*1024),b''): hasher.update(data)
        record={'source':str(relative),'project':relative.parts[0] if len(relative.parts)>1 else source.name,
                'bytes':path.stat().st_size,'sha256':hasher.hexdigest()}
        old=existing.get(str(relative))
        if old and old.get('status')=='converted' and old.get('sha256')==record['sha256'] and path.suffix.lower()!='.csv' and (output/old.get('textFile','_missing')).is_file():
            records.append(old);continue
        try:
            text,warnings=convert(path)
            if not text.strip(): raise ValueError('没有提取到文字')
            target=output/'texts'/relative.parent/(relative.name+'.txt');target.parent.mkdir(parents=True,exist_ok=True)
            header=f'原文件：{relative}\n原始 SHA-256：{record["sha256"]}\n项目目录：{record["project"]}\n'
            target.write_text(header+'\n'+text+'\n',encoding='utf-8')
            record.update(status='converted',textFile=str(target.relative_to(output)),characters=len(text),warnings=warnings)
        except Exception as error: record.update(status='not_converted',reason=str(error))
        records.append(record)
        if len(records)%10==0: print(json.dumps({'processed':len(records),'file':str(relative)},ensure_ascii=False),flush=True)
    manifest={'format':'mockinterview-preprocessed-v1','createdAt':datetime.now(timezone.utc).isoformat(),
              'sourceRoot':str(source),'records':records}
    (output/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    projects={}
    for row in records:
        item=projects.setdefault(row['project'],{'converted':0,'other':0,'characters':0})
        item['converted' if row['status']=='converted' else 'other']+=1;item['characters']+=row.get('characters',0)
    lines=['# 面试资料预处理索引','', '原文件未修改；本目录仅包含文字提取结果，不代表事实已经核实。',
           '代码保留原始行号与函数边界，文档保留段落、页/幻灯片和工作表定位。',
           '图片、视频、权重、扫描页及图中机制不能从文字提取中推断。','',
           '## 建题优先顺序','',
           '先从简历中的实际项目与个人贡献定位，再关注目标、选型、关键机制、结果、失败和局限。',
           '报告/论文/最终汇报用于概览；关键源码用于实现追问；辅助脚本与原始数据不直接逐文件出题。',
           '候选项目、优化建议和草稿必须区分于已经完成的经历；“预计常问”不等于统计高频。','',
           '| 项目目录 | 已转文字 | 未转文件 | 文字量 |','|---|---:|---:|---:|']
    for name,row in projects.items(): lines.append(f'| {name} | {row["converted"]} | {row["other"]} | {row["characters"]} |')
    lines+=['','## 提取异常与范围限制','']
    for row in records:
        if row['status']=='not_converted': lines.append(f'- {row["source"]}：{row["reason"]}')
        else:
            for warning in row['warnings']: lines.append(f'- {row["source"]}：{warning}')
    (output/'README.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    print(json.dumps({'output':str(output),'converted':sum(r['status']=='converted' for r in records),
                      'other':sum(r['status']!='converted' for r in records),'projects':projects},ensure_ascii=False))

if __name__=='__main__': main()
